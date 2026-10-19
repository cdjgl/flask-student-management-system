"""学生管理系统 - Flask 主应用（MySQL 版）。

数据存储：MySQL（PyMySQL + DBUtils 连接池），配置见 config.py / .env。
"""
import os
import json
import zipfile
import io
from datetime import datetime, timedelta
from decimal import Decimal
from functools import wraps

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from werkzeug.security import generate_password_hash, check_password_hash
from flask import (Flask, render_template, request, redirect, url_for,
                   session, jsonify, make_response, flash)
from urllib.parse import quote

import config
import db

app = Flask(__name__)
app.secret_key = config.SECRET_KEY
app.permanent_session_lifetime = timedelta(hours=config.SESSION_LIFETIME_HOURS)

BASE_DIR = os.path.dirname(__file__)
BACKUPS_DIR = os.path.join(BASE_DIR, 'backups')
EXPORTS_DIR = os.path.join(BASE_DIR, 'exports')
os.makedirs(BACKUPS_DIR, exist_ok=True)
os.makedirs(EXPORTS_DIR, exist_ok=True)

# 启动时建表并确保默认管理员存在
db.init_db()

# ---- 列别名：DB 列 class_name 在模板中仍以 `class` 呈现，保持模板零改动 ----
STUDENT_COLS = 'student_id, name, gender, college, class_name AS `class`'
DORM_COLS = 'student_id, name, college, class_name AS `class`, building, room, bed'
COMP_COLS = ('certificate_id, competition_name, award_level, award_year, '
             'college, class_name AS `class`, award_students, contact, advisor')

# 备份/恢复支持的表：表名 -> (列清单, 导出 SQL)
BACKUP_TABLES = {
    'users': ('username, password_hash, role, email, contact, register_time, last_login_time',
              "SELECT username, password_hash, role, email, contact, "
              "DATE_FORMAT(register_time, '%Y-%m-%d %H:%i:%s') AS register_time, "
              "IFNULL(DATE_FORMAT(last_login_time, '%Y-%m-%d %H:%i:%s'), '') AS last_login_time FROM users"),
    'classes': ('class_id, college, class_name, advisor, created_at',
                'SELECT class_id, college, class_name, advisor, created_at FROM classes'),
    'students': ('student_id, name, gender, college, class_name',
                 'SELECT student_id, name, gender, college, class_name FROM students'),
    'courses': ('student_id, course_name, credit, score, course_type, course_nature, semester, gpa, credit_gpa',
                'SELECT student_id, course_name, credit, score, course_type, course_nature, '
                'semester, gpa, credit_gpa FROM courses'),
    'dormitories': ('student_id, name, college, class_name, building, room, bed',
                    'SELECT student_id, name, college, class_name, building, room, bed FROM dormitories'),
    'competitions': ('certificate_id, competition_name, award_level, award_year, college, '
                     'class_name, award_students, contact, advisor',
                     'SELECT certificate_id, competition_name, award_level, award_year, college, '
                     'class_name, award_students, contact, advisor FROM competitions'),
    'notices': ('notice_id, title, content, priority, publisher, publish_time',
                "SELECT notice_id, title, content, priority, publisher, "
                "DATE_FORMAT(publish_time, '%Y-%m-%d %H:%i:%s') AS publish_time FROM notices"),
    'warning_status': ('student_id, warning_type, warning_key, status',
                       'SELECT student_id, warning_type, warning_key, status FROM warning_status'),
}


def _f(value):
    """Decimal -> float，便于模板渲染与 JSON 序列化。"""
    return float(value) if isinstance(value, Decimal) else value


def normalize_course(row):
    for key in ('credit', 'score', 'gpa', 'credit_gpa'):
        if key in row:
            row[key] = _f(row[key])
    return row


def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'username' not in session:
            return redirect(url_for('login', next=request.url))
        return f(*args, **kwargs)
    return decorated_function


def calculate_gpa(score):
    if score >= 90:
        return 5.0 - (100 - score) * 0.1
    elif score >= 60:
        return 1.0 + (score - 60) * 0.1
    else:
        return 0.0


def build_pages(total_pages, page):
    """复刻原分页窗口逻辑。"""
    if total_pages <= 10:
        return list(range(1, total_pages + 1))
    if page <= 5:
        return list(range(1, 7)) + ['...'] + [total_pages]
    if page >= total_pages - 4:
        return [1] + ['...'] + list(range(total_pages - 5, total_pages + 1))
    return [1] + ['...'] + list(range(page - 2, page + 3)) + ['...'] + [total_pages]


def style_header(ws, start_color='4096ff'):
    for cell in ws[1]:
        cell.font = Font(bold=True, color='ffffff')
        cell.alignment = Alignment(horizontal='center')
        cell.fill = PatternFill(start_color=start_color, end_color=start_color, fill_type='solid')


def send_xlsx(wb, filename):
    filepath = os.path.join(EXPORTS_DIR, filename)
    wb.save(filepath)
    with open(filepath, 'rb') as f:
        data = f.read()
    response = make_response(data)
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response.headers['Content-Disposition'] = f"attachment; filename*=UTF-8''{quote(filename)}"
    return response


def get_courses_by_student(student_ids):
    """一次取多名学生的课程并按学生分组，课程保持插入顺序。"""
    grouped = {sid: [] for sid in student_ids}
    if not student_ids:
        return grouped
    placeholders = ', '.join(['%s'] * len(student_ids))
    rows = db.query(
        f'SELECT * FROM courses WHERE student_id IN ({placeholders}) ORDER BY id',
        tuple(student_ids))
    for row in rows:
        grouped[row['student_id']].append(normalize_course(row))
    return grouped


def generate_warnings():
    """实时计算成绩预警（挂科 + 低绩点），并合并 warning_status 中的处理状态。

    警情 ID 形如 `学号|类型|键`，可稳定定位到具体课程，用于状态持久化。
    """
    students_list = db.query(f'SELECT {STUDENT_COLS} FROM students ORDER BY id')
    courses_grouped = get_courses_by_student([s['student_id'] for s in students_list])
    status_rows = db.query('SELECT student_id, warning_type, warning_key, status FROM warning_status')
    status_map = {(r['student_id'], r['warning_type'], r['warning_key']): r['status']
                  for r in status_rows}

    warnings = []
    for student in students_list:
        student_id = student['student_id']
        courses = courses_grouped.get(student_id, [])

        for course in courses:
            score = _f(course.get('score', 0))
            if score < 60:
                key = course.get('course_name', '未知课程')
                warnings.append({
                    'warning_id': f'{student_id}|fail|{key}',
                    'student_id': student_id,
                    'name': student['name'],
                    'college': student.get('college', ''),
                    'class_name': student.get('class', ''),
                    'warning_type': 'fail',
                    'reason': f'课程「{key}」成绩不及格',
                    'score': f'{score:g}分',
                    'status': status_map.get((student_id, 'fail', key), 'pending'),
                })

        if courses:
            total_credits = sum(_f(c.get('credit', 0)) for c in courses)
            total_gpa_points = sum(_f(c.get('credit', 0)) * calculate_gpa(_f(c.get('score', 0)))
                                   for c in courses)
            avg_gpa = total_gpa_points / total_credits if total_credits > 0 else 0
            if avg_gpa < 1.5:
                warnings.append({
                    'warning_id': f'{student_id}|low_gpa|avg',
                    'student_id': student_id,
                    'name': student['name'],
                    'college': student.get('college', ''),
                    'class_name': student.get('class', ''),
                    'warning_type': 'low_gpa',
                    'reason': f'平均绩点 {avg_gpa:.2f} < 1.5',
                    'score': f'{avg_gpa:.2f}',
                    'status': status_map.get((student_id, 'low_gpa', 'avg'), 'pending'),
                })
    return warnings


def save_warning_status(warning_id, status):
    """按警情 ID 持久化处理状态。"""
    try:
        student_id, warning_type, warning_key = warning_id.split('|', 2)
    except ValueError:
        return
    db.execute(
        'INSERT INTO warning_status (student_id, warning_type, warning_key, status) '
        'VALUES (%s, %s, %s, %s) '
        'ON DUPLICATE KEY UPDATE status = VALUES(status)',
        (student_id, warning_type, warning_key, status))


@app.route('/')
def index():
    if 'username' in session:
        return redirect(url_for('dashboard'))
    return redirect(url_for('login'))


@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        remember = request.form.get('remember') == 'on'

        user = db.query_one('SELECT * FROM users WHERE username = %s', (username,))
        if user and check_password_hash(user['password_hash'], password):
            session['username'] = username
            if remember:
                session.permanent = True
            db.execute('UPDATE users SET last_login_time = NOW() WHERE username = %s', (username,))
            return redirect(url_for('dashboard'))

        return render_template('login.html', error='用户名或密码错误')

    return render_template('login.html')


@app.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        email = request.form.get('email', '').strip()
        password = request.form.get('password', '')
        confirm_password = request.form.get('confirm_password', '')

        if not username or not email or not password:
            return render_template('register.html', error='所有字段都是必填的')
        if password != confirm_password:
            return render_template('register.html', error='两次输入的密码不一致')
        if len(password) < 6:
            return render_template('register.html', error='密码长度至少为6位')

        if db.query_one('SELECT id FROM users WHERE username = %s', (username,)):
            return render_template('register.html', error='用户名已存在')
        if db.query_one('SELECT id FROM users WHERE email = %s', (email,)):
            return render_template('register.html', error='邮箱已被注册')

        db.execute(
            'INSERT INTO users (username, password_hash, role, email, register_time) '
            'VALUES (%s, %s, %s, %s, NOW())',
            (username, generate_password_hash(password), 'user', email))

        return redirect(url_for('login', registered='true'))

    return render_template('register.html')


@app.route('/logout')
def logout():
    session.pop('username', None)
    return redirect(url_for('login'))


@app.route('/dashboard')
@login_required
def dashboard():
    student_count = db.query_one('SELECT COUNT(*) AS c FROM students')['c']
    course_count = db.query_one('SELECT COUNT(*) AS c FROM courses')['c']
    dormitory_count = db.query_one('SELECT COUNT(*) AS c FROM dormitories')['c']
    competition_count = db.query_one('SELECT COUNT(*) AS c FROM competitions')['c']

    return render_template('dashboard.html',
                           username=session['username'],
                           student_count=student_count,
                           course_count=course_count,
                           dormitory_count=dormitory_count,
                           competition_count=competition_count,
                           current_time=datetime.now().strftime('%Y年%m月%d日 %H:%M:%S'))


@app.route('/students')
@login_required
def students():
    search = request.args.get('search', '').strip()
    college = request.args.get('college', '')
    class_name = request.args.get('class', '')
    gender = request.args.get('gender', '')
    page = int(request.args.get('page', 1))
    per_page = 8

    where, params = [], []
    if search:
        where.append('(student_id LIKE %s OR name LIKE %s)')
        params += [f'%{search}%', f'%{search}%']
    if college:
        where.append('college = %s')
        params.append(college)
    if class_name:
        where.append('class_name = %s')
        params.append(class_name)
    if gender:
        where.append('gender = %s')
        params.append(gender)
    where_sql = ('WHERE ' + ' AND '.join(where)) if where else ''

    total = db.query_one(f'SELECT COUNT(*) AS c FROM students {where_sql}', tuple(params))['c']
    total_pages = (total + per_page - 1) // per_page
    page = max(1, min(page, max(total_pages, 1)))
    offset = (page - 1) * per_page

    rows = db.query(
        f'SELECT {STUDENT_COLS} FROM students {where_sql} '
        f'ORDER BY student_id LIMIT %s OFFSET %s',
        tuple(params + [per_page, offset]))
    courses_grouped = get_courses_by_student([r['student_id'] for r in rows])
    for row in rows:
        row['courses'] = courses_grouped.get(row['student_id'], [])

    colleges = [r['college'] for r in db.query(
        'SELECT DISTINCT college FROM students WHERE college <> \'\' ORDER BY college')]
    classes = [r['class'] for r in db.query(
        'SELECT DISTINCT class_name AS `class` FROM students WHERE class_name <> \'\' ORDER BY class_name')]

    return render_template('students.html',
                           students=rows,
                           colleges=colleges,
                           classes=classes,
                           search=search,
                           selected_college=college,
                           selected_class=class_name,
                           selected_gender=gender,
                           current_page=page,
                           total_pages=total_pages,
                           total=total,
                           pages=build_pages(total_pages, page))


@app.route('/student/add', methods=['POST'])
@login_required
def add_student():
    student_id = request.form.get('student_id', '').strip()
    name = request.form.get('name', '').strip()
    gender = request.form.get('gender', '男')
    college = request.form.get('college', '').strip()
    class_name = request.form.get('class', '').strip()

    if not student_id or not name:
        return redirect(url_for('students', error='学号和姓名为必填项'))
    if db.query_one('SELECT id FROM students WHERE student_id = %s', (student_id,)):
        return redirect(url_for('students', error='学号已存在'))

    db.execute(
        'INSERT INTO students (student_id, name, gender, college, class_name) '
        'VALUES (%s, %s, %s, %s, %s)',
        (student_id, name, gender, college, class_name))
    return redirect(url_for('students'))


@app.route('/student/<student_id>/edit', methods=['POST'])
@login_required
def edit_student(student_id):
    name = request.form.get('name', '').strip()
    gender = request.form.get('gender', '男')
    college = request.form.get('college', '').strip()
    class_name = request.form.get('class', '').strip()

    db.execute(
        'UPDATE students SET name = %s, gender = %s, college = %s, class_name = %s '
        'WHERE student_id = %s',
        (name, gender, college, class_name, student_id))
    return redirect(url_for('students'))


@app.route('/student/<student_id>/delete', methods=['POST'])
@login_required
def delete_student(student_id):
    db.execute_in_transaction([
        ('DELETE FROM dormitories WHERE student_id = %s', (student_id,)),
        ('DELETE FROM students WHERE student_id = %s', (student_id,)),
    ])
    return redirect(url_for('students'))


@app.route('/student/<student_id>/courses')
@login_required
def student_courses(student_id):
    student = db.query_one(
        f'SELECT {STUDENT_COLS} FROM students WHERE student_id = %s', (student_id,))
    if not student:
        return redirect(url_for('students'))

    student['courses'] = [normalize_course(r) for r in
                          db.query('SELECT * FROM courses WHERE student_id = %s ORDER BY id',
                                   (student_id,))]
    return render_template('student_courses.html', student=student)


@app.route('/student/<student_id>/course/add', methods=['POST'])
@login_required
def add_course(student_id):
    course_name = request.form.get('course_name', '').strip()
    credit = float(request.form.get('credit', 0))
    score = int(request.form.get('score', 0))
    course_type = request.form.get('course_type', '专业核心课')
    course_nature = request.form.get('course_nature', '必修')
    semester = request.form.get('semester', '2024-2025第一学期')

    gpa = calculate_gpa(score)
    credit_gpa = credit * gpa

    db.execute(
        'INSERT INTO courses (student_id, course_name, credit, score, course_type, '
        'course_nature, semester, gpa, credit_gpa) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)',
        (student_id, course_name, credit, score, course_type, course_nature,
         semester, round(gpa, 2), round(credit_gpa, 2)))
    return redirect(url_for('student_courses', student_id=student_id))


def _course_id_by_index(student_id, course_index):
    row = db.query_one(
        'SELECT id FROM courses WHERE student_id = %s ORDER BY id LIMIT 1 OFFSET %s',
        (student_id, course_index))
    return row['id'] if row else None


@app.route('/student/<student_id>/course/<int:course_index>/edit', methods=['POST'])
@login_required
def edit_course(student_id, course_index):
    course_id = _course_id_by_index(student_id, course_index)
    if course_id is None:
        return redirect(url_for('student_courses', student_id=student_id))

    course_name = request.form.get('course_name', '').strip()
    credit = float(request.form.get('credit', 0))
    score = int(request.form.get('score', 0))
    course_type = request.form.get('course_type', '专业核心课')
    course_nature = request.form.get('course_nature', '必修')
    semester = request.form.get('semester', '2024-2025第一学期')

    gpa = calculate_gpa(score)
    credit_gpa = credit * gpa

    db.execute(
        'UPDATE courses SET course_name = %s, credit = %s, score = %s, course_type = %s, '
        'course_nature = %s, semester = %s, gpa = %s, credit_gpa = %s WHERE id = %s',
        (course_name, credit, score, course_type, course_nature,
         semester, round(gpa, 2), round(credit_gpa, 2), course_id))
    return redirect(url_for('student_courses', student_id=student_id))


@app.route('/student/<student_id>/course/<int:course_index>/delete', methods=['POST'])
@login_required
def delete_course(student_id, course_index):
    course_id = _course_id_by_index(student_id, course_index)
    if course_id is not None:
        db.execute('DELETE FROM courses WHERE id = %s', (course_id,))
    return redirect(url_for('student_courses', student_id=student_id))


@app.route('/dormitories')
@login_required
def dormitories():
    college = request.args.get('college', '')
    class_name = request.args.get('class', '')
    building = request.args.get('building', '')
    room = request.args.get('room', '')
    bed = request.args.get('bed', '')
    page = int(request.args.get('page', 1))
    per_page = 8

    where, params = [], []
    for column, value in (('college', college), ('class_name', class_name),
                          ('building', building), ('room', room), ('bed', bed)):
        if value:
            where.append(f'{column} = %s')
            params.append(value)
    where_sql = ('WHERE ' + ' AND '.join(where)) if where else ''

    total = db.query_one(f'SELECT COUNT(*) AS c FROM dormitories {where_sql}', tuple(params))['c']
    total_pages = (total + per_page - 1) // per_page
    page = max(1, min(page, max(total_pages, 1)))
    offset = (page - 1) * per_page

    rows = db.query(
        f'SELECT {DORM_COLS} FROM dormitories {where_sql} '
        f'ORDER BY student_id LIMIT %s OFFSET %s',
        tuple(params + [per_page, offset]))

    def distinct(column):
        return [r['v'] for r in db.query(
            f'SELECT DISTINCT {column} AS v FROM dormitories WHERE {column} <> \'\' ORDER BY v')]

    return render_template('dormitories.html',
                           dormitories=rows,
                           colleges=distinct('college'),
                           classes=distinct('class_name'),
                           buildings=distinct('building'),
                           rooms=distinct('room'),
                           beds=distinct('bed'),
                           selected_college=college,
                           selected_class=class_name,
                           selected_building=building,
                           selected_room=room,
                           selected_bed=bed,
                           current_page=page,
                           total_pages=total_pages,
                           total=total,
                           pages=build_pages(total_pages, page))


@app.route('/dormitory/add', methods=['POST'])
@login_required
def add_dormitory():
    student_id = request.form.get('student_id', '').strip()
    name = request.form.get('name', '').strip()
    college = request.form.get('college', '').strip()
    class_name = request.form.get('class', '').strip()
    building = request.form.get('building', '').strip()
    room = request.form.get('room', '').strip()
    bed = request.form.get('bed', '').strip()

    if not student_id or not name:
        return redirect(url_for('dormitories', error='学号和姓名为必填项'))
    if db.query_one('SELECT id FROM dormitories WHERE student_id = %s', (student_id,)):
        return redirect(url_for('dormitories', error='该学号已存在宿舍信息'))

    db.execute(
        'INSERT INTO dormitories (student_id, name, college, class_name, building, room, bed) '
        'VALUES (%s, %s, %s, %s, %s, %s, %s)',
        (student_id, name, college, class_name, building, room, bed))
    return redirect(url_for('dormitories'))


@app.route('/dormitory/<student_id>/edit', methods=['POST'])
@login_required
def edit_dormitory(student_id):
    name = request.form.get('name', '').strip()
    college = request.form.get('college', '').strip()
    class_name = request.form.get('class', '').strip()
    building = request.form.get('building', '').strip()
    room = request.form.get('room', '').strip()
    bed = request.form.get('bed', '').strip()

    db.execute(
        'UPDATE dormitories SET name = %s, college = %s, class_name = %s, '
        'building = %s, room = %s, bed = %s WHERE student_id = %s',
        (name, college, class_name, building, room, bed, student_id))
    return redirect(url_for('dormitories'))


@app.route('/dormitory/<student_id>/delete', methods=['POST'])
@login_required
def delete_dormitory(student_id):
    db.execute('DELETE FROM dormitories WHERE student_id = %s', (student_id,))
    return redirect(url_for('dormitories'))


@app.route('/competitions')
@login_required
def competitions():
    college = request.args.get('college', '')
    class_name = request.args.get('class', '')
    award_level = request.args.get('award_level', '')
    award_year = request.args.get('award_year', '')
    page = int(request.args.get('page', 1))
    per_page = 8

    where, params = [], []
    for column, value in (('college', college), ('class_name', class_name),
                          ('award_level', award_level), ('award_year', award_year)):
        if value:
            where.append(f'{column} = %s')
            params.append(value)
    where_sql = ('WHERE ' + ' AND '.join(where)) if where else ''

    total = db.query_one(f'SELECT COUNT(*) AS c FROM competitions {where_sql}', tuple(params))['c']
    total_pages = (total + per_page - 1) // per_page
    page = max(1, min(page, max(total_pages, 1)))
    offset = (page - 1) * per_page

    rows = db.query(
        f'SELECT {COMP_COLS} FROM competitions {where_sql} '
        f'ORDER BY certificate_id LIMIT %s OFFSET %s',
        tuple(params + [per_page, offset]))

    def distinct(column):
        return [r['v'] for r in db.query(
            f'SELECT DISTINCT {column} AS v FROM competitions WHERE {column} <> \'\' ORDER BY v')]

    years = [r['v'] for r in db.query(
        'SELECT DISTINCT award_year AS v FROM competitions WHERE award_year <> \'\' ORDER BY v DESC')]

    return render_template('competitions.html',
                           competitions=rows,
                           colleges=distinct('college'),
                           classes=distinct('class_name'),
                           levels=distinct('award_level'),
                           years=years,
                           selected_college=college,
                           selected_class=class_name,
                           selected_level=award_level,
                           selected_year=award_year,
                           current_page=page,
                           total_pages=total_pages,
                           total=total,
                           pages=build_pages(total_pages, page))


@app.route('/competition/add', methods=['POST'])
@login_required
def add_competition():
    certificate_id = request.form.get('certificate_id', '').strip()
    competition_name = request.form.get('competition_name', '').strip()
    award_level = request.form.get('award_level', '校级')
    award_year = request.form.get('award_year', '').strip()
    college = request.form.get('college', '').strip()
    class_name = request.form.get('class', '').strip()
    award_students = request.form.get('award_students', '').strip()
    contact = request.form.get('contact', '').strip()
    advisor = request.form.get('advisor', '').strip()

    if not certificate_id or not competition_name:
        return redirect(url_for('competitions', error='证书编号和竞赛名称为必填项'))
    if db.query_one('SELECT id FROM competitions WHERE certificate_id = %s', (certificate_id,)):
        return redirect(url_for('competitions', error='证书编号已存在'))

    db.execute(
        'INSERT INTO competitions (certificate_id, competition_name, award_level, award_year, '
        'college, class_name, award_students, contact, advisor) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)',
        (certificate_id, competition_name, award_level, award_year, college,
         class_name, award_students, contact, advisor))
    return redirect(url_for('competitions'))


@app.route('/competition/<certificate_id>/edit', methods=['POST'])
@login_required
def edit_competition(certificate_id):
    competition_name = request.form.get('competition_name', '').strip()
    award_level = request.form.get('award_level', '校级')
    award_year = request.form.get('award_year', '').strip()
    college = request.form.get('college', '').strip()
    class_name = request.form.get('class', '').strip()
    award_students = request.form.get('award_students', '').strip()
    contact = request.form.get('contact', '').strip()
    advisor = request.form.get('advisor', '').strip()

    db.execute(
        'UPDATE competitions SET competition_name = %s, award_level = %s, award_year = %s, '
        'college = %s, class_name = %s, award_students = %s, contact = %s, advisor = %s '
        'WHERE certificate_id = %s',
        (competition_name, award_level, award_year, college, class_name,
         award_students, contact, advisor, certificate_id))
    return redirect(url_for('competitions'))


@app.route('/competition/<certificate_id>/delete', methods=['POST'])
@login_required
def delete_competition(certificate_id):
    db.execute('DELETE FROM competitions WHERE certificate_id = %s', (certificate_id,))
    return redirect(url_for('competitions'))


@app.route('/classes')
@login_required
def classes():
    college = request.args.get('college', '')
    search_name = request.args.get('search_name', '')
    advisor = request.args.get('advisor', '')
    page = int(request.args.get('page', 1))
    per_page = 10

    where, params = [], []
    if college:
        where.append('c.college = %s')
        params.append(college)
    if search_name:
        where.append('c.class_name LIKE %s')
        params.append(f'%{search_name}%')
    if advisor:
        where.append('c.advisor = %s')
        params.append(advisor)
    where_sql = ('WHERE ' + ' AND '.join(where)) if where else ''

    total = db.query_one(
        f'SELECT COUNT(*) AS c FROM classes c {where_sql}', tuple(params))['c']
    total_pages = (total + per_page - 1) // per_page
    page = max(1, min(page, max(total_pages, 1)))
    offset = (page - 1) * per_page

    rows = db.query(
        'SELECT c.class_id, c.college, c.class_name, c.advisor, c.created_at, '
        '(SELECT COUNT(*) FROM students s WHERE s.class_name = c.class_name) AS student_count '
        f'FROM classes c {where_sql} ORDER BY c.class_id LIMIT %s OFFSET %s',
        tuple(params + [per_page, offset]))

    colleges = [r['college'] for r in db.query(
        'SELECT DISTINCT college FROM classes WHERE college <> \'\' ORDER BY college')]
    advisors = [r['advisor'] for r in db.query(
        'SELECT DISTINCT advisor FROM classes WHERE advisor <> \'\' ORDER BY advisor')]

    return render_template('classes.html',
                           classes=rows,
                           colleges=colleges,
                           advisors=advisors,
                           selected_college=college,
                           selected_search_name=search_name,
                           selected_advisor=advisor,
                           current_page=page,
                           total_pages=total_pages,
                           total=total,
                           pages=build_pages(total_pages, page))


@app.route('/class/add', methods=['POST'])
@login_required
def add_class():
    college = request.form.get('college', '').strip()
    class_name = request.form.get('class_name', '').strip()
    advisor = request.form.get('advisor', '').strip()

    if not college or not class_name:
        return redirect(url_for('classes', error='学院和班级名称都是必填的'))
    if db.query_one('SELECT id FROM classes WHERE class_name = %s', (class_name,)):
        return redirect(url_for('classes', error='班级名称已存在'))

    max_num = db.query_one(
        "SELECT COALESCE(MAX(CAST(SUBSTRING(class_id, 6) AS UNSIGNED)), 0) AS m "
        "FROM classes WHERE class_id LIKE 'CLASS%'")['m']
    class_id = f'CLASS{int(max_num or 0) + 1:03d}'
    created_at = datetime.now().strftime('%Y-%m-%d')

    db.execute(
        'INSERT INTO classes (class_id, college, class_name, advisor, created_at) '
        'VALUES (%s, %s, %s, %s, %s)',
        (class_id, college, class_name, advisor, created_at))
    return redirect(url_for('classes'))


@app.route('/class/<class_id>/edit', methods=['POST'])
@login_required
def edit_class(class_id):
    college = request.form.get('college', '').strip()
    class_name = request.form.get('class_name', '').strip()
    advisor = request.form.get('advisor', '').strip()

    if not college or not class_name:
        return redirect(url_for('classes', error='学院和班级名称都是必填的'))

    target = db.query_one('SELECT class_name FROM classes WHERE class_id = %s', (class_id,))
    if not target:
        return redirect(url_for('classes', error='班级不存在'))

    exists = db.query_one(
        'SELECT id FROM classes WHERE class_id <> %s AND class_name = %s',
        (class_id, class_name))
    if exists:
        return redirect(url_for('classes', error='班级名称已存在'))

    old_class_name = target['class_name']
    operations = [
        ('UPDATE classes SET college = %s, class_name = %s, advisor = %s WHERE class_id = %s',
         (college, class_name, advisor, class_id)),
    ]
    if old_class_name != class_name:
        operations += [
            ('UPDATE students SET class_name = %s WHERE class_name = %s',
             (class_name, old_class_name)),
            ('UPDATE dormitories SET class_name = %s WHERE class_name = %s',
             (class_name, old_class_name)),
            ('UPDATE competitions SET class_name = %s WHERE class_name = %s',
             (class_name, old_class_name)),
        ]
    db.execute_in_transaction(operations)

    return redirect(url_for('classes'))


@app.route('/class/<class_id>/delete', methods=['POST'])
@login_required
def delete_class(class_id):
    target = db.query_one(
        'SELECT class_name, (SELECT COUNT(*) FROM students s WHERE s.class_name = classes.class_name) '
        'AS student_count FROM classes WHERE class_id = %s', (class_id,))
    if not target:
        return redirect(url_for('classes', error='班级不存在'))
    if target['student_count'] > 0:
        return redirect(url_for('classes', error='该班级还有学生，无法删除'))

    db.execute('DELETE FROM classes WHERE class_id = %s', (class_id,))
    return redirect(url_for('classes'))


@app.route('/export/classes')
@login_required
def export_classes():
    college = request.args.get('college', '')
    search_name = request.args.get('search_name', '')
    advisor = request.args.get('advisor', '')

    where, params = [], []
    if college:
        where.append('c.college = %s')
        params.append(college)
    if search_name:
        where.append('c.class_name LIKE %s')
        params.append(f'%{search_name}%')
    if advisor:
        where.append('c.advisor = %s')
        params.append(advisor)
    where_sql = ('WHERE ' + ' AND '.join(where)) if where else ''

    rows = db.query(
        'SELECT c.class_id, c.college, c.class_name, c.advisor, c.created_at, '
        '(SELECT COUNT(*) FROM students s WHERE s.class_name = c.class_name) AS student_count '
        f'FROM classes c {where_sql} ORDER BY c.class_id', tuple(params))

    wb = Workbook()
    ws = wb.active
    ws.title = '班级信息'
    ws.append(['班级ID', '二级学院', '班级名称', '辅导员', '学生人数', '创建时间'])
    style_header(ws)
    for cls in rows:
        ws.append([cls['class_id'], cls['college'], cls['class_name'],
                   cls['advisor'], cls['student_count'], cls['created_at']])

    filename = f'班级信息_{len(rows)}条.xlsx'
    return send_xlsx(wb, filename)


@app.route('/grade-warning')
@login_required
def grade_warning():
    warning_type = request.args.get('warning_type', '')
    college = request.args.get('college', '')
    class_name = request.args.get('class', '')
    status = request.args.get('status', '')
    page = int(request.args.get('page', 1))
    per_page = 10

    warnings = generate_warnings()
    filtered_warnings = [w for w in warnings
                         if (not warning_type or w['warning_type'] == warning_type)
                         and (not college or w['college'] == college)
                         and (not class_name or w['class_name'] == class_name)
                         and (not status or w['status'] == status)]

    total = len(filtered_warnings)
    total_pages = (total + per_page - 1) // per_page
    page = max(1, min(page, max(total_pages, 1)))
    start = (page - 1) * per_page
    paginated_warnings = filtered_warnings[start:start + per_page]

    colleges = [r['college'] for r in db.query(
        'SELECT DISTINCT college FROM students WHERE college <> \'\' ORDER BY college')]
    classes = [r['class'] for r in db.query(
        'SELECT DISTINCT class_name AS `class` FROM students WHERE class_name <> \'\' ORDER BY class_name')]

    return render_template('grade_warning.html',
                           warnings=paginated_warnings,
                           colleges=colleges,
                           classes=classes,
                           selected_warning_type=warning_type,
                           selected_college=college,
                           selected_class=class_name,
                           selected_status=status,
                           current_page=page,
                           total_pages=total_pages,
                           total=total,
                           pages=build_pages(total_pages, page))


@app.route('/update-warning-status', methods=['POST'])
@login_required
def update_warning_status():
    warning_id = request.form.get('warning_id', '')
    status = request.form.get('status', '')
    save_warning_status(warning_id, status)
    return jsonify({'success': True})


@app.route('/batch-resolve-warnings', methods=['POST'])
@login_required
def batch_resolve_warnings():
    warning_ids = request.form.getlist('warning_ids')
    for warning_id in warning_ids:
        save_warning_status(warning_id, 'resolved')
    return redirect(url_for('grade_warning'))


@app.route('/export/grade-warning')
@login_required
def export_grade_warning():
    warning_type = request.args.get('warning_type', '')
    college = request.args.get('college', '')
    class_name = request.args.get('class', '')
    status = request.args.get('status', '')

    warnings = generate_warnings()
    filtered_warnings = [w for w in warnings
                         if (not warning_type or w['warning_type'] == warning_type)
                         and (not college or w['college'] == college)
                         and (not class_name or w['class_name'] == class_name)
                         and (not status or w['status'] == status)]

    wb = Workbook()
    ws = wb.active
    ws.title = '成绩预警报表'
    ws.append(['预警ID', '学号', '姓名', '二级学院', '班级', '预警类型', '预警原因', '成绩/绩点', '处理状态'])
    style_header(ws)

    status_mapping = {'pending': '未处理', 'contacted': '已沟通', 'resolved': '已解决'}
    type_mapping = {'fail': '挂科预警', 'low_gpa': '低绩点预警'}

    for warning in filtered_warnings:
        ws.append([
            warning['warning_id'],
            warning['student_id'],
            warning['name'],
            warning['college'],
            warning['class_name'],
            type_mapping.get(warning['warning_type'], warning['warning_type']),
            warning['reason'],
            warning['score'],
            status_mapping.get(warning['status'], warning['status']),
        ])

    today = datetime.now().strftime('%Y%m%d')
    filename = f'成绩预警报表_{today}.xlsx'
    return send_xlsx(wb, filename)


@app.route('/statistics')
@login_required
def statistics():
    student_count = db.query_one('SELECT COUNT(*) AS c FROM students')['c']
    course_count = db.query_one('SELECT COUNT(*) AS c FROM courses')['c']
    dormitory_count = db.query_one('SELECT COUNT(*) AS c FROM dormitories')['c']
    competition_count = db.query_one('SELECT COUNT(*) AS c FROM competitions')['c']

    gender_rows = db.query(
        'SELECT gender, COUNT(*) AS c FROM students GROUP BY gender')
    male_count = next((r['c'] for r in gender_rows if r['gender'] == '男'), 0)
    female_count = next((r['c'] for r in gender_rows if r['gender'] == '女'), 0)
    gender_distribution = []
    if student_count > 0:
        gender_distribution = [
            {'gender': '男', 'count': male_count, 'percentage': round(male_count / student_count * 100, 2)},
            {'gender': '女', 'count': female_count, 'percentage': round(female_count / student_count * 100, 2)},
        ]

    range_defs = [
        ('0-59分（不及格）', 0, 59), ('60-69分（及格）', 60, 69), ('70-79分（中等）', 70, 79),
        ('80-89分（良好）', 80, 89), ('90-100分（优秀）', 90, 100),
    ]
    score_ranges = []
    total_courses = course_count
    for label, low, high in range_defs:
        count = db.query_one(
            'SELECT COUNT(*) AS c FROM courses WHERE score >= %s AND score <= %s',
            (low, high))['c']
        score_ranges.append({
            'range': label, 'min': low, 'max': high, 'count': count,
            'percentage': round(count / total_courses * 100, 2) if total_courses > 0 else 0,
        })

    course_rows = db.query(
        'SELECT course_name, AVG(score) AS avg_score, MAX(score) AS max_score, '
        'MIN(score) AS min_score, SUM(score >= 60) AS pass_count, COUNT(*) AS total '
        'FROM courses GROUP BY course_name ORDER BY avg_score DESC')
    course_list = [{
        'name': r['course_name'],
        'avg_score': round(_f(r['avg_score']), 2),
        'max_score': _f(r['max_score']),
        'min_score': _f(r['min_score']),
        'pass_rate': round(r['pass_count'] / r['total'] * 100, 2),
        'student_count': r['total'],
    } for r in course_rows]

    return render_template('statistics.html',
                           student_count=student_count,
                           course_count=course_count,
                           dormitory_count=dormitory_count,
                           competition_count=competition_count,
                           gender_distribution=gender_distribution,
                           score_ranges=score_ranges,
                           course_list=course_list[:10])


@app.route('/visualization')
@login_required
def visualization():
    return render_template('visualization.html')


@app.route('/api/dashboard/stats')
@login_required
def api_dashboard_stats():
    total_students = db.query_one('SELECT COUNT(*) AS c FROM students')['c']
    total_courses = db.query_one('SELECT COUNT(*) AS c FROM courses')['c']
    total_competitions = db.query_one('SELECT COUNT(*) AS c FROM competitions')['c']
    college_count = db.query_one(
        'SELECT COUNT(DISTINCT college) AS c FROM students WHERE college <> \'\'')['c']
    warning_count = len(generate_warnings())

    return jsonify({
        'total_students': total_students,
        'total_courses': total_courses,
        'total_competitions': total_competitions,
        'warning_count': warning_count,
        'college_count': college_count,
    })


@app.route('/api/chart/gender')
@login_required
def api_chart_gender():
    rows = db.query('SELECT gender, COUNT(*) AS c FROM students GROUP BY gender')
    male = next((r['c'] for r in rows if r['gender'] == '男'), 0)
    female = next((r['c'] for r in rows if r['gender'] == '女'), 0)
    return jsonify({'labels': ['男生', '女生'], 'data': [male, female]})


@app.route('/api/chart/score-distribution')
@login_required
def api_chart_score_distribution():
    row = db.query_one(
        'SELECT SUM(score < 60) AS a, SUM(score >= 60 AND score < 70) AS b, '
        'SUM(score >= 70 AND score < 80) AS c, SUM(score >= 80 AND score < 90) AS d, '
        'SUM(score >= 90) AS e FROM courses')
    data = [_f(row[k] or 0) for k in ('a', 'b', 'c', 'd', 'e')]
    return jsonify({'labels': ['0-59', '60-69', '70-79', '80-89', '90-100'], 'data': data})


@app.route('/api/chart/college-distribution')
@login_required
def api_chart_college_distribution():
    rows = db.query(
        'SELECT college, COUNT(*) AS c FROM students GROUP BY college '
        'ORDER BY c DESC LIMIT 10')
    return jsonify({'labels': [r['college'] for r in rows], 'data': [r['c'] for r in rows]})


@app.route('/api/chart/pass-rate')
@login_required
def api_chart_pass_rate():
    row = db.query_one(
        'SELECT COUNT(*) AS total, SUM(score >= 60) AS passed FROM courses')
    total = row['total']
    passed = int(row['passed'] or 0)
    rate = round(passed / total * 100, 2) if total > 0 else 0
    return jsonify({'pass_rate': rate, 'total': total, 'passed': passed})


@app.route('/api/chart/course-averages')
@login_required
def api_chart_course_averages():
    rows = db.query(
        'SELECT course_name, AVG(score) AS avg_score FROM courses '
        'GROUP BY course_name ORDER BY avg_score DESC LIMIT 10')
    return jsonify({'labels': [r['course_name'] for r in rows],
                    'data': [round(_f(r['avg_score']), 2) for r in rows]})


@app.route('/api/chart/class-distribution')
@login_required
def api_chart_class_distribution():
    rows = db.query(
        'SELECT class_name, COUNT(*) AS c FROM students GROUP BY class_name ORDER BY c DESC')
    return jsonify({'labels': [r['class_name'] for r in rows], 'data': [r['c'] for r in rows]})


@app.route('/api/chart/gpa-distribution')
@login_required
def api_chart_gpa_distribution():
    # 与原逻辑一致：按学生取课程平均绩点（对 gpa 字段求平均）分箱
    rows = db.query(
        'SELECT student_id, AVG(gpa) AS avg_gpa FROM courses GROUP BY student_id')
    ranges = {'0-1': 0, '1-2': 0, '2-3': 0, '3-4': 0, '4-5': 0}
    for row in rows:
        avg_gpa = _f(row['avg_gpa'])
        if avg_gpa < 1:
            ranges['0-1'] += 1
        elif avg_gpa < 2:
            ranges['1-2'] += 1
        elif avg_gpa < 3:
            ranges['2-3'] += 1
        elif avg_gpa < 4:
            ranges['3-4'] += 1
        else:
            ranges['4-5'] += 1
    return jsonify({'labels': list(ranges.keys()), 'data': list(ranges.values())})


@app.route('/api/chart/course-type')
@login_required
def api_chart_course_type():
    rows = db.query(
        'SELECT course_type, COUNT(*) AS c FROM courses GROUP BY course_type')
    return jsonify({'labels': [r['course_type'] for r in rows], 'data': [r['c'] for r in rows]})


@app.route('/api/chart/semester')
@login_required
def api_chart_semester():
    rows = db.query(
        'SELECT semester, COUNT(*) AS c FROM courses GROUP BY semester ORDER BY semester')
    return jsonify({'labels': [r['semester'] for r in rows], 'data': [r['c'] for r in rows]})


@app.route('/api/chart/course-nature')
@login_required
def api_chart_course_nature():
    rows = db.query(
        'SELECT course_nature, COUNT(*) AS c FROM courses GROUP BY course_nature')
    return jsonify({'labels': [r['course_nature'] for r in rows], 'data': [r['c'] for r in rows]})


@app.route('/api/chart/summary')
@login_required
def api_chart_summary():
    row = db.query_one(
        'SELECT COUNT(*) AS total, SUM(score >= 60) AS passed, SUM(score >= 90) AS high_score, '
        'AVG(score) AS avg_score, AVG(gpa) AS avg_gpa FROM courses')
    total = row['total']
    if total > 0:
        passed = int(row['passed'] or 0)
        high_score = int(row['high_score'] or 0)
        return jsonify({
            'pass_rate': round(passed / total * 100, 2),
            'avg_score': round(_f(row['avg_score']), 1),
            'avg_gpa': round(_f(row['avg_gpa']), 2),
            'high_score_rate': round(high_score / total * 100, 2),
        })
    return jsonify({'pass_rate': 0, 'avg_score': 0, 'avg_gpa': 0, 'high_score_rate': 0})


@app.route('/export/students')
@login_required
def export_students():
    search = request.args.get('search', '')
    college = request.args.get('college', '')
    class_name = request.args.get('class', '')
    gender = request.args.get('gender', '')

    where, params = [], []
    if search:
        where.append('(student_id LIKE %s OR name LIKE %s)')
        params += [f'%{search}%', f'%{search}%']
    if college:
        where.append('college = %s')
        params.append(college)
    if class_name:
        where.append('class_name = %s')
        params.append(class_name)
    if gender:
        where.append('gender = %s')
        params.append(gender)
    where_sql = ('WHERE ' + ' AND '.join(where)) if where else ''

    rows = db.query(
        'SELECT s.student_id, s.name, s.gender, s.college, s.class_name, '
        'COUNT(c.id) AS course_count, COALESCE(SUM(c.credit), 0) AS total_credit, '
        'COALESCE(SUM(c.credit_gpa) / NULLIF(SUM(c.credit), 0), 0) AS avg_gpa '
        f'FROM students s LEFT JOIN courses c ON c.student_id = s.student_id {where_sql} '
        'GROUP BY s.student_id, s.name, s.gender, s.college, s.class_name ORDER BY s.student_id',
        tuple(params))

    wb = Workbook()
    ws = wb.active
    ws.title = '学生信息'
    headers = ['学号', '姓名', '性别', '二级学院', '班级', '课程数', '总学分', '平均绩点']
    ws.append(headers)
    style_header(ws)
    for row in rows:
        ws.cell(row=ws.max_row + 1, column=1, value=row['student_id'])
        ws.cell(row=ws.max_row, column=2, value=row['name'])
        ws.cell(row=ws.max_row, column=3, value=row['gender'])
        ws.cell(row=ws.max_row, column=4, value=row['college'])
        ws.cell(row=ws.max_row, column=5, value=row['class_name'])
        ws.cell(row=ws.max_row, column=6, value=row['course_count'])
        ws.cell(row=ws.max_row, column=7, value=_f(row['total_credit']))
        ws.cell(row=ws.max_row, column=8, value=round(_f(row['avg_gpa']), 2))

    filename = f'学生信息_{len(rows)}条.xlsx'
    return send_xlsx(wb, filename)


@app.route('/export/dormitories')
@login_required
def export_dormitories():
    college = request.args.get('college', '')
    class_name = request.args.get('class', '')
    building = request.args.get('building', '')
    room = request.args.get('room', '')
    bed = request.args.get('bed', '')

    where, params = [], []
    for column, value in (('college', college), ('class_name', class_name),
                          ('building', building), ('room', room), ('bed', bed)):
        if value:
            where.append(f'{column} = %s')
            params.append(value)
    where_sql = ('WHERE ' + ' AND '.join(where)) if where else ''

    rows = db.query(
        f'SELECT {DORM_COLS} FROM dormitories {where_sql} ORDER BY student_id', tuple(params))

    wb = Workbook()
    ws = wb.active
    ws.title = '宿舍信息'
    ws.append(['学号', '姓名', '二级学院', '班级', '宿舍楼', '房间号', '床位号'])
    style_header(ws)
    for dorm in rows:
        ws.append([dorm['student_id'], dorm['name'], dorm['college'], dorm['class'],
                   dorm['building'], dorm['room'], dorm['bed']])

    filename = f'宿舍信息_{len(rows)}条.xlsx'
    return send_xlsx(wb, filename)


@app.route('/export/competitions')
@login_required
def export_competitions():
    college = request.args.get('college', '')
    class_name = request.args.get('class', '')
    award_level = request.args.get('award_level', '')
    award_year = request.args.get('award_year', '')

    where, params = [], []
    for column, value in (('college', college), ('class_name', class_name),
                          ('award_level', award_level), ('award_year', award_year)):
        if value:
            where.append(f'{column} = %s')
            params.append(value)
    where_sql = ('WHERE ' + ' AND '.join(where)) if where else ''

    rows = db.query(
        f'SELECT {COMP_COLS} FROM competitions {where_sql} ORDER BY certificate_id', tuple(params))

    wb = Workbook()
    ws = wb.active
    ws.title = '竞赛信息'
    ws.append(['证书编号', '竞赛名称', '获奖级别', '获奖年份', '二级学院', '班级', '获奖同学', '联系方式', '指导老师'])
    style_header(ws)
    for comp in rows:
        ws.append([comp['certificate_id'], comp['competition_name'], comp['award_level'],
                   comp['award_year'], comp['college'], comp['class'],
                   comp['award_students'], comp['contact'], comp['advisor']])

    filename = f'竞赛信息_{len(rows)}条.xlsx'
    return send_xlsx(wb, filename)


@app.route('/export/statistics')
@login_required
def export_statistics():
    student_count = db.query_one('SELECT COUNT(*) AS c FROM students')['c']
    course_count = db.query_one('SELECT COUNT(*) AS c FROM courses')['c']
    dormitory_count = db.query_one('SELECT COUNT(*) AS c FROM dormitories')['c']
    competition_count = db.query_one('SELECT COUNT(*) AS c FROM competitions')['c']

    wb = Workbook()
    ws1 = wb.active
    ws1.title = '概况统计'
    ws1.append(['学生总数', student_count])
    ws1.append(['课程总数', course_count])
    ws1.append(['宿舍信息数', dormitory_count])
    ws1.append(['竞赛获奖数', competition_count])

    ws2 = wb.create_sheet('性别分布')
    ws2.append(['性别', '人数', '比例(%)'])
    gender_rows = db.query('SELECT gender, COUNT(*) AS c FROM students GROUP BY gender')
    for row in gender_rows:
        percentage = round(row['c'] / student_count * 100, 2) if student_count > 0 else 0
        ws2.append([row['gender'], row['c'], percentage])

    ws3 = wb.create_sheet('成绩分布')
    ws3.append(['分数段', '课程数', '比例(%)'])
    range_defs = [
        ('0-59分（不及格）', 0, 59), ('60-69分（及格）', 60, 69), ('70-79分（中等）', 70, 79),
        ('80-89分（良好）', 80, 89), ('90-100分（优秀）', 90, 100),
    ]
    for label, low, high in range_defs:
        count = db.query_one(
            'SELECT COUNT(*) AS c FROM courses WHERE score >= %s AND score <= %s',
            (low, high))['c']
        percentage = round(count / course_count * 100, 2) if course_count > 0 else 0
        ws3.append([label, count, percentage])

    ws4 = wb.create_sheet('课程统计')
    ws4.append(['课程名称', '平均分', '最高分', '最低分', '及格率(%)', '选课人数'])
    course_rows = db.query(
        'SELECT course_name, AVG(score) AS avg_score, MAX(score) AS max_score, '
        'MIN(score) AS min_score, SUM(score >= 60) AS pass_count, COUNT(*) AS total '
        'FROM courses GROUP BY course_name ORDER BY avg_score DESC')
    for row in course_rows:
        ws4.append([
            row['course_name'], round(_f(row['avg_score']), 2), _f(row['max_score']),
            _f(row['min_score']), round(row['pass_count'] / row['total'] * 100, 2), row['total'],
        ])

    filename = '统计分析报表.xlsx'
    return send_xlsx(wb, filename)


@app.route('/notices')
@login_required
def notices():
    priority = request.args.get('priority', '')
    start_date = request.args.get('start_date', '')
    end_date = request.args.get('end_date', '')
    keyword = request.args.get('keyword', '')
    page = int(request.args.get('page', 1))

    where, params = [], []
    if priority:
        where.append('priority = %s')
        params.append(priority)
    if start_date:
        where.append('publish_time >= %s')
        params.append(start_date)
    if end_date:
        where.append('publish_time < DATE_ADD(%s, INTERVAL 1 DAY)')
        params.append(end_date)
    if keyword:
        where.append('title LIKE %s')
        params.append(f'%{keyword}%')
    where_sql = ('WHERE ' + ' AND '.join(where)) if where else ''

    total = db.query_one(f'SELECT COUNT(*) AS c FROM notices {where_sql}', tuple(params))['c']
    per_page = 10
    total_pages = (total + per_page - 1) // per_page
    page = max(1, min(page, max(total_pages, 1)))
    offset = (page - 1) * per_page

    rows = db.query(
        "SELECT notice_id, title, priority, publisher, "
        "DATE_FORMAT(publish_time, '%%Y-%%m-%%d %%H:%%i:%%s') AS publish_time "
        f'FROM notices {where_sql} '
        "ORDER BY (priority = 'urgent') DESC, publish_time DESC LIMIT %s OFFSET %s",
        tuple(params + [per_page, offset]))

    return render_template('notices.html',
                           notices=rows,
                           priority=priority,
                           start_date=start_date,
                           end_date=end_date,
                           keyword=keyword,
                           page=page,
                           total_pages=total_pages)


@app.route('/notices/view/<notice_id>')
@login_required
def view_notice(notice_id):
    notice = db.query_one(
        "SELECT notice_id, title, content, priority, publisher, "
        "DATE_FORMAT(publish_time, '%%Y-%%m-%%d %%H:%%i:%%s') AS publish_time "
        'FROM notices WHERE notice_id = %s', (notice_id,))
    if not notice:
        return render_template('404.html'), 404
    return render_template('notice_detail.html', notice=notice)


@app.route('/notices/add', methods=['GET', 'POST'])
@login_required
def add_notice():
    if session.get('username') != 'admin':
        return redirect(url_for('notices'))

    if request.method == 'POST':
        title = request.form.get('title', '').strip()
        content = request.form.get('content', '').strip()
        priority = request.form.get('priority', 'normal')

        if not title or not content:
            flash('标题和内容不能为空', 'danger')
            return redirect(url_for('add_notice'))

        max_num = db.query_one(
            "SELECT COALESCE(MAX(CAST(SUBSTRING(notice_id, 7) AS UNSIGNED)), 0) AS m "
            "FROM notices WHERE notice_id LIKE 'NOTICE%'")['m']
        new_id = f'NOTICE{int(max_num or 0) + 1:03d}'

        db.execute(
            'INSERT INTO notices (notice_id, title, content, priority, publisher, publish_time) '
            'VALUES (%s, %s, %s, %s, %s, NOW())',
            (new_id, title, content, priority, session.get('username', 'admin')))

        flash('公告发布成功', 'success')
        return redirect(url_for('notices'))

    return render_template('notice_form.html',
                           title='发布公告',
                           action_url=url_for('add_notice'),
                           notice=None)


@app.route('/notices/edit/<notice_id>', methods=['GET', 'POST'])
@login_required
def edit_notice(notice_id):
    if session.get('username') != 'admin':
        return redirect(url_for('notices'))

    notice = db.query_one('SELECT * FROM notices WHERE notice_id = %s', (notice_id,))
    if not notice:
        return render_template('404.html'), 404
    notice['publish_time'] = notice['publish_time'].strftime('%Y-%m-%d %H:%M:%S')

    if request.method == 'POST':
        title = request.form.get('title', '').strip()
        content = request.form.get('content', '').strip()
        priority = request.form.get('priority', 'normal')

        if not title or not content:
            flash('标题和内容不能为空', 'danger')
            return redirect(url_for('edit_notice', notice_id=notice_id))

        db.execute(
            'UPDATE notices SET title = %s, content = %s, priority = %s WHERE notice_id = %s',
            (title, content, priority, notice_id))

        flash('公告编辑成功', 'success')
        return redirect(url_for('notices'))

    return render_template('notice_form.html',
                           title='编辑公告',
                           action_url=url_for('edit_notice', notice_id=notice_id),
                           notice=notice)


@app.route('/notices/delete/<notice_id>')
@login_required
def delete_notice(notice_id):
    if session.get('username') != 'admin':
        return redirect(url_for('notices'))

    if db.execute('DELETE FROM notices WHERE notice_id = %s', (notice_id,)):
        flash('公告删除成功', 'success')
    return redirect(url_for('notices'))


def _current_user():
    return db.query_one(
        "SELECT username, role, email, contact, "
        "DATE_FORMAT(register_time, '%%Y-%%m-%%d %%H:%%i:%%s') AS register_time, "
        "IFNULL(DATE_FORMAT(last_login_time, '%%Y-%%m-%%d %%H:%%i:%%s'), '从未登录') AS last_login_time "
        'FROM users WHERE username = %s', (session.get('username'),))


@app.route('/profile')
@login_required
def profile():
    return render_template('profile.html',
                           user=_current_user(),
                           old_password_error='',
                           password_mismatch='',
                           email_error='')


@app.route('/profile/change-password', methods=['POST'])
@login_required
def change_password():
    old_password = request.form.get('old_password', '')
    new_password = request.form.get('new_password', '')
    confirm_password = request.form.get('confirm_password', '')

    user = _current_user()
    row = db.query_one('SELECT password_hash FROM users WHERE username = %s',
                       (session.get('username'),))

    if not row or not check_password_hash(row['password_hash'], old_password):
        return render_template('profile.html', user=user,
                               old_password_error='原密码错误',
                               password_mismatch='', email_error='')
    if new_password != confirm_password:
        return render_template('profile.html', user=user,
                               old_password_error='',
                               password_mismatch='两次输入的密码不一致', email_error='')
    if len(new_password) < 6:
        return render_template('profile.html', user=user,
                               old_password_error='',
                               password_mismatch='密码长度至少6位', email_error='')

    db.execute('UPDATE users SET password_hash = %s WHERE username = %s',
               (generate_password_hash(new_password), session.get('username')))

    session.clear()
    flash('密码修改成功，请重新登录', 'success')
    return redirect(url_for('login'))


@app.route('/profile/edit', methods=['POST'])
@login_required
def edit_profile():
    email = request.form.get('email', '').strip()
    contact = request.form.get('contact', '').strip()

    user = _current_user()
    username = session.get('username')

    if email:
        if db.query_one('SELECT id FROM users WHERE username <> %s AND email = %s',
                        (username, email)):
            return render_template('profile.html', user=user,
                                   old_password_error='', password_mismatch='',
                                   email_error='该邮箱已被其他用户使用')
        if '@' not in email:
            return render_template('profile.html', user=user,
                                   old_password_error='', password_mismatch='',
                                   email_error='请输入有效的邮箱地址')
        db.execute('UPDATE users SET email = %s WHERE username = %s', (email, username))

    db.execute('UPDATE users SET contact = %s WHERE username = %s', (contact, username))

    flash('个人信息修改成功', 'success')
    return redirect(url_for('profile'))


@app.route('/backup')
@login_required
def backup():
    if session.get('username') != 'admin':
        return redirect(url_for('dashboard'))

    page = int(request.args.get('page', 1))
    per_page = 10

    backups = []
    for filename in os.listdir(BACKUPS_DIR):
        if filename.endswith('.zip'):
            filepath = os.path.join(BACKUPS_DIR, filename)
            if os.path.isfile(filepath):
                file_size = os.path.getsize(filepath)
                backup_time = os.path.getmtime(filepath)
                backups.append({
                    'filename': filename,
                    'backup_time': datetime.fromtimestamp(backup_time).strftime('%Y-%m-%d %H:%M:%S'),
                    'size': format_file_size(file_size),
                })

    backups.sort(key=lambda x: x['backup_time'], reverse=True)

    total = len(backups)
    total_pages = (total + per_page - 1) // per_page
    page = max(1, min(page, max(total_pages, 1)))
    start = (page - 1) * per_page
    paginated_backups = backups[start:start + per_page]

    return render_template('backup.html',
                           backups=paginated_backups,
                           page=page,
                           total_pages=total_pages)


def format_file_size(size):
    if size < 1024:
        return f'{size} B'
    elif size < 1024 * 1024:
        return f'{size / 1024:.2f} KB'
    else:
        return f'{size / (1024 * 1024):.2f} MB'


def _dump_table(table, select_sql):
    rows = db.query(select_sql)
    for row in rows:
        for key, value in row.items():
            if isinstance(value, Decimal):
                row[key] = float(value)
    return rows


@app.route('/backup/create')
@login_required
def create_backup():
    if session.get('username') != 'admin':
        return redirect(url_for('dashboard'))

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    backup_filename = f'系统备份_{timestamp}.zip'
    backup_path = os.path.join(BACKUPS_DIR, backup_filename)

    with zipfile.ZipFile(backup_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
        for table, (_, select_sql) in BACKUP_TABLES.items():
            rows = _dump_table(table, select_sql)
            content = json.dumps({table: rows}, ensure_ascii=False, indent=2, default=str)
            zipf.writestr(f'{table}.json', content)

    flash('数据备份成功', 'success')
    return redirect(url_for('backup'))


@app.route('/backup/download/<filename>')
@login_required
def download_backup(filename):
    if session.get('username') != 'admin':
        return redirect(url_for('dashboard'))

    filepath = os.path.join(BACKUPS_DIR, filename)
    if not os.path.exists(filepath):
        flash('备份文件不存在', 'danger')
        return redirect(url_for('backup'))

    with open(filepath, 'rb') as f:
        data = f.read()

    response = make_response(data)
    response.headers['Content-Type'] = 'application/zip'
    response.headers['Content-Disposition'] = f"attachment; filename*=UTF-8''{quote(filename)}"
    return response


@app.route('/backup/restore/<filename>')
@login_required
def restore_backup(filename):
    if session.get('username') != 'admin':
        return redirect(url_for('dashboard'))

    backup_path = os.path.join(BACKUPS_DIR, filename)
    if not os.path.exists(backup_path):
        flash('备份文件不存在', 'danger')
        return redirect(url_for('backup'))

    try:
        with zipfile.ZipFile(backup_path, 'r') as zipf:
            payloads = {}
            for table in BACKUP_TABLES:
                if f'{table}.json' in zipf.namelist():
                    data = json.loads(zipf.read(f'{table}.json').decode('utf-8'))
                    payloads[table] = data.get(table, [])

        db.clear_tables(list(BACKUP_TABLES))
        for table in BACKUP_TABLES:
            rows = payloads.get(table) or []
            if not rows:
                continue
            columns, _ = BACKUP_TABLES[table]
            col_list = columns.split(', ')
            insert_sql = (f'INSERT INTO {table} ({columns}) '
                          f"VALUES ({', '.join(['%s'] * len(col_list))})")
            db.execute_many(insert_sql, [tuple(row[col] for col in col_list) for row in rows])

        flash('数据恢复成功', 'success')
    except Exception as e:
        flash(f'数据恢复失败：{str(e)}', 'danger')

    return redirect(url_for('backup'))


@app.route('/backup/delete/<filename>')
@login_required
def delete_backup(filename):
    if session.get('username') != 'admin':
        return redirect(url_for('dashboard'))

    filepath = os.path.join(BACKUPS_DIR, filename)
    if os.path.exists(filepath):
        os.remove(filepath)
        flash('备份文件已删除', 'success')
    else:
        flash('备份文件不存在', 'danger')

    return redirect(url_for('backup'))


@app.errorhandler(404)
def page_not_found(e):
    return render_template('404.html'), 404


@app.errorhandler(500)
def internal_server_error(e):
    return render_template('500.html'), 500


if __name__ == '__main__':
    # 开发模式下排除备份/导出目录，避免文件写入触发热重载中断请求
    app.run(debug=True, host='0.0.0.0', port=5000,
            exclude_patterns=['backups/*', 'exports/*'])
